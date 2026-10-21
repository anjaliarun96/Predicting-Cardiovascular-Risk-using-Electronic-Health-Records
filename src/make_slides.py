"""Generate the review slide deck (.pptx) from results/*. Run AFTER train.py.

    python src/make_slides.py --results results --out slides/deck.pptx --names "Name1, Name2"
"""
import argparse
import json
import os

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results")
    ap.add_argument("--out", default="slides/deck.pptx")
    ap.add_argument("--names", default="<Team members>")
    a = ap.parse_args()
    R = a.results
    ds = json.load(open(f"{R}/data_summary.json"))
    m = pd.read_csv(f"{R}/metrics.csv")
    sub = pd.read_csv(f"{R}/subgroups.csv")
    best = m.sort_values("test_auroc").iloc[-1]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def bullets(title, items, notes=None):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        body = sl.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5)
        tf = body.text_frame
        tf.word_wrap = True
        for i, t in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = t
            p.font.size = Pt(22)
        if notes:
            sl.notes_slide.notes_text_frame.text = notes
        return sl

    def picture_slide(title, path, caption=None):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = title
        if os.path.exists(path):
            sl.shapes.add_picture(path, Inches(3.4), Inches(1.4), height=Inches(5.3))
        if caption:
            tb = sl.shapes.add_textbox(Inches(0.5), Inches(6.8), Inches(12), Inches(0.5))
            tb.text_frame.text = caption
        return sl

    sl = prs.slides.add_slide(prs.slide_layouts[0])
    sl.shapes.title.text = "Predicting Cardiovascular Risk using Electronic Health Records"
    sl.placeholders[1].text = f"UE24CS352A Machine Learning - Mini-Project\n{a.names}"

    bullets("Problem statement", [
        "CVD is the leading cause of death globally - early identification of at-risk patients matters",
        "Task: predict presence of cardiovascular disease from routine clinical measurements",
        "Goal: compare interpretable regularised logistic regression vs gradient boosting (XGBoost)",
        "Inspired by Nguyen (Stanford CS229) EHR-based CVD risk prediction; STRIDE data is private, so a public dataset is used"])
    bullets("Dataset", [
        f"Kaggle Cardiovascular Disease dataset: {ds['n_raw']:,} records -> {ds['n_clean']:,} after cleaning",
        "Features: age, gender, height, weight, BP (systolic/diastolic), cholesterol, glucose, smoking, alcohol, activity",
        "Engineered: BMI, pulse pressure",
        f"Target balanced-ish: {ds['positive_rate']*100:.1f}% positive",
        "Cleaning: implausible BP/height/weight removed, duplicates dropped",
        "80/20 stratified train/test split; tuning with cross-validation on train only"])
    bullets("Approach", [
        "Baselines: Logistic Regression with Ridge / Lasso / Elastic Net (standardised features)",
        "Main model: XGBoost - sequential regularised decision trees, each correcting the previous errors",
        f"Hyperparameters tuned by {ds['cv_folds']}-fold CV grid search on AUROC",
        "Metrics: AUROC (primary), accuracy, sensitivity, specificity",
        "Fairness check: AUROC by gender and age group"])
    bullets("Implementation", [
        "src/data.py - load, clean, feature engineering",
        "src/train.py - tuning, evaluation, plots, subgroup analysis",
        "src/make_report.py / make_slides.py - build write-up and slides from results",
        "Python: pandas, scikit-learn, XGBoost, matplotlib"])

    mt = m.copy()
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = "Results"
    rows, cols = len(mt) + 1, 5
    tbl = sl.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.6) * rows).table
    for j, h in enumerate(["Model", "CV AUROC", "Test AUROC", "Sensitivity", "Specificity"]):
        tbl.cell(0, j).text = h
    for i, r in enumerate(mt.itertuples(), start=1):
        for j, v in enumerate([r.model, f"{r.cv_auroc:.3f}", f"{r.test_auroc:.3f}", f"{r.sensitivity:.3f}", f"{r.specificity:.3f}"]):
            tbl.cell(i, j).text = v
    picture_slide("ROC curves (test set)", f"{R}/roc.png")
    picture_slide("What drives the predictions?", f"{R}/feature_importance.png", "XGBoost feature importance")
    bullets(f"Subgroup (fairness) check - {best.model}",
            [f"{r.subgroup}: AUROC {r.auroc:.3f} (n={r.n})" for r in sub.itertuples()] +
            ["Performance is not identical across groups -> audit before clinical use"])
    bullets("Conclusions & future work", [
        f"Best model: {best.model}, test AUROC {best.test_auroc:.3f}; small CV-test gap -> little over-fitting",
        "Regularisation helps linear models; boosting captures non-linear effects",
        "Limitations: prevalent (not incident) CVD, self-reported fields, no ethnicity data",
        "Future: calibration, longitudinal EHR features, fairness-aware training"])
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
